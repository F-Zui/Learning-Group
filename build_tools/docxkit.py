# -*- coding: utf-8 -*-
"""docxkit —— 生成算法竞赛模板 docx 的排版工具层。

提供标题 / 正文 / 列表 / 代码块 / 提示框 / 表格 等原语，
以及统一的中文字体、行距、页面设置。
"""
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---------------------------------------------------------------- 颜色主题
C_MAIN = "1F3864"      # 主标题深蓝
C_SUB = "2E5496"       # 二级标题蓝
C_SUB3 = "8C4B00"      # 三级标题棕（算法名）
C_CODE = "1A1A1A"      # 代码文字
C_CODE_BG = "F4F5F7"   # 代码底色
C_HEAD_BG = "1F3864"   # 表头底色
C_WARN_BG = "FDF0E3"   # 提示框底色
C_TIP_BG = "EAF1FB"    # 技巧框底色
C_ALT_BG = "F2F5FA"    # 表格隔行


def _set_run(run, name="Times New Roman", cjk="微软雅黑", size=10.5,
             bold=False, color=None, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:ascii"), name)
    rf.set(qn("w:hAnsi"), name)
    rf.set(qn("w:eastAsia"), cjk)
    return run


def shade(element, fill):
    """给段落或单元格加底色。"""
    pr = element._element.get_or_add_pPr() if hasattr(element, "paragraph_format") \
        else element._element.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    pr.append(sh)


def cell_shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    tcpr.append(sh)


def para_shade(p, fill):
    ppr = p._element.get_or_add_pPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    ppr.append(sh)


def para_border(p, left=None, box=None, size=18):
    ppr = p._element.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    sides = ["top", "left", "bottom", "right"] if box else ["left"]
    for side in sides:
        el = OxmlElement("w:" + side)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), box or left)
        bd.append(el)
    ppr.append(bd)


def keep_with_next(p):
    ppr = p._element.get_or_add_pPr()
    el = OxmlElement("w:keepNext")
    ppr.append(el)


class Builder:
    def __init__(self):
        self.doc = Document()
        self._setup()
        self._toc = []

    # ------------------------------------------------------------ 页面设置
    def _setup(self):
        st = self.doc.styles["Normal"]
        st.font.name = "Times New Roman"
        st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        pf = st.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.28
        pf.space_after = Pt(3)

        # 覆盖内置 Heading 样式：目录域依赖这些样式，但外观完全由我们决定
        for name, size, color, before, after in (
                ("Heading 1", 19, C_MAIN, 20, 9),
                ("Heading 2", 14.5, C_SUB, 13, 5),
                ("Heading 3", 12.5, C_SUB3, 10, 3)):
            s = self.doc.styles[name]
            s.font.name = "Times New Roman"
            s.font.size = Pt(size)
            s.font.bold = True
            s.font.color.rgb = RGBColor.from_string(color)
            rpr = s.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rpr.append(rf)
            rf.set(qn("w:ascii"), "Times New Roman")
            rf.set(qn("w:hAnsi"), "Times New Roman")
            rf.set(qn("w:eastAsia"), "微软雅黑")
            spf = s.paragraph_format
            spf.space_before = Pt(before)
            spf.space_after = Pt(after)
            spf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            spf.line_spacing = 1.2
            spf.keep_with_next = True

        sec = self.doc.sections[0]
        sec.top_margin = Cm(2.2)
        sec.bottom_margin = Cm(2.0)
        sec.left_margin = Cm(2.3)
        sec.right_margin = Cm(2.0)

    # ------------------------------------------------------------ 基本元素
    def _hstyled(self, level, text):
        """用 Word 内置 Heading 样式创建段落（这样目录域才能识别），字体与颜色自行覆盖。"""
        p = self.doc.add_paragraph(style="Heading %d" % level)
        p.add_run(text)
        return p

    def title(self, text, sub=None):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        _set_run(p.add_run(text), cjk="微软雅黑", size=26, bold=True, color=C_MAIN)
        if sub:
            q = self.doc.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.space_after = Pt(14)
            _set_run(q.add_run(sub), cjk="微软雅黑", size=11.5, color="595959")

    def h1(self, text, toc=True):
        p = self._hstyled(1, text)
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(9)
        p.runs[0].text = text
        _set_run(p.runs[0], cjk="微软雅黑", size=19, bold=True, color=C_MAIN)
        para_border(p, left=C_MAIN, size=30)
        para_shade(p, "EDF2FA")
        keep_with_next(p)
        if toc:
            self._toc.append((1, text))
        return p

    def h2(self, text, toc=True):
        p = self._hstyled(2, text)
        p.paragraph_format.space_before = Pt(13)
        p.paragraph_format.space_after = Pt(5)
        _set_run(p.runs[0], cjk="微软雅黑", size=14.5, bold=True, color=C_SUB)
        keep_with_next(p)
        if toc:
            self._toc.append((2, text))
        return p

    def h3(self, text, tag=None, toc=True):
        p = self._hstyled(3, "▍" + text)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        _set_run(p.runs[0], cjk="微软雅黑", size=12.5, bold=True, color=C_SUB3)
        if tag:
            _set_run(p.add_run("  " + tag), cjk="微软雅黑", size=9, color="7F7F7F")
        keep_with_next(p)
        if toc:
            self._toc.append((3, text))
        return p

    def p(self, *segs, indent=0.0, space_after=3.5):
        """segs 里可以混入 (text, style_dict)。"""
        pa = self.doc.add_paragraph()
        pa.paragraph_format.space_after = Pt(space_after)
        if indent:
            pa.paragraph_format.left_indent = Cm(indent)
        for seg in segs:
            if isinstance(seg, tuple):
                txt, st = seg
                _set_run(pa.add_run(txt), cjk=st.get("cjk", "微软雅黑"),
                         name=st.get("name", "Times New Roman"),
                         size=st.get("size", 10.5), bold=st.get("bold", False),
                         color=st.get("color"), italic=st.get("italic", False))
            else:
                _set_run(pa.add_run(seg), cjk="微软雅黑", size=10.5)
        return pa

    def kv(self, key, value, indent=0.0):
        """「用途：xxx」这种标签行。"""
        return self.p((key, {"bold": True, "color": C_SUB}), value, indent=indent)

    def bullets(self, items, indent=0.6, mark="·"):
        for it in items:
            pa = self.doc.add_paragraph()
            pa.paragraph_format.left_indent = Cm(indent + 0.35)
            pa.paragraph_format.first_line_indent = Cm(-0.35)
            pa.paragraph_format.space_after = Pt(2)
            _set_run(pa.add_run(mark + " "), cjk="微软雅黑", size=10.5, bold=True, color=C_SUB)
            if isinstance(it, tuple):
                for seg in it:
                    if isinstance(seg, tuple):
                        txt, st = seg
                        _set_run(pa.add_run(txt), cjk=st.get("cjk", "微软雅黑"),
                                 name=st.get("name", "Times New Roman"),
                                 size=st.get("size", 10.5), bold=st.get("bold", False),
                                 color=st.get("color"))
                    else:
                        _set_run(pa.add_run(seg), cjk="微软雅黑", size=10.5)
            else:
                _set_run(pa.add_run(it), cjk="微软雅黑", size=10.5)

    def code(self, text, caption=None, size=8.8):
        if caption:
            cp = self.doc.add_paragraph()
            cp.paragraph_format.space_before = Pt(4)
            cp.paragraph_format.space_after = Pt(1)
            _set_run(cp.add_run(caption), cjk="微软雅黑", size=9, bold=True, color="595959")
            keep_with_next(cp)
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        cell_shade(cell, C_CODE_BG)
        cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
        lines = text.strip("\n").split("\n")
        for i, ln in enumerate(lines):
            pa = cell.add_paragraph()
            pa.paragraph_format.space_after = Pt(0)
            pa.paragraph_format.space_before = Pt(0)
            pa.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            pa.paragraph_format.line_spacing = 1.05
            run = pa.add_run(ln if ln.strip() else " ")
            _set_run(run, name="Consolas", cjk="Microsoft YaHei", size=size, color=C_CODE)
        self._tbl_borders(t, width=6)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(0)
        sp.paragraph_format.space_before = Pt(0)
        for r in sp.runs:
            pass
        _set_run(sp.add_run(""), size=4)
        return t

    def _tbl_borders(self, table, width=6, color="D0D5DD"):
        tblpr = table._element.tblPr
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement("w:" + side)
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(width))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
            borders.append(el)
        tblpr.append(borders)

    def box(self, title, lines, kind="tip"):
        """提示框：kind = tip / warn / ok"""
        fill = {"tip": C_TIP_BG, "warn": C_WARN_BG, "ok": "EAF6EC"}[kind]
        border = {"tip": "9DC3F0", "warn": "E8B77A", "ok": "9BD3A6"}[kind]
        t = self.doc.add_table(rows=1, cols=1)
        cell = t.rows[0].cells[0]
        cell_shade(cell, fill)
        p0 = cell.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        _set_run(p0.add_run(title), cjk="微软雅黑", size=10.5, bold=True, color="333333")
        for ln in lines:
            pa = cell.add_paragraph()
            pa.paragraph_format.space_after = Pt(1)
            pa.paragraph_format.left_indent = Cm(0.25)
            if isinstance(ln, tuple):
                for seg in ln:
                    if isinstance(seg, tuple):
                        txt, st = seg
                        _set_run(pa.add_run(txt), cjk=st.get("cjk", "微软雅黑"),
                                 name=st.get("name", "Times New Roman"),
                                 size=st.get("size", 10), bold=st.get("bold", False),
                                 color=st.get("color"))
                    else:
                        _set_run(pa.add_run(seg), cjk="微软雅黑", size=10)
            else:
                _set_run(pa.add_run(ln), cjk="微软雅黑", size=10)
        self._tbl_borders(t, width=4, color=border)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(0)
        _set_run(sp.add_run(""), size=4)
        return t

    def table(self, header, rows, widths=None, size=9.5, header_size=9.5):
        t = self.doc.add_table(rows=1, cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            cell_shade(c, C_HEAD_BG)
            pa = c.paragraphs[0]
            pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pa.paragraph_format.space_after = Pt(1)
            pa.paragraph_format.space_before = Pt(1)
            _set_run(pa.add_run(h), cjk="微软雅黑", size=header_size, bold=True, color="FFFFFF")
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                c = cells[i]
                if ri % 2 == 1:
                    cell_shade(c, C_ALT_BG)
                pa = c.paragraphs[0]
                pa.paragraph_format.space_after = Pt(1)
                pa.paragraph_format.space_before = Pt(1)
                if isinstance(v, tuple):
                    txt, st = v
                    _set_run(pa.add_run(txt), cjk="微软雅黑", name=st.get("name", "Times New Roman"),
                             size=st.get("size", size), bold=st.get("bold", False),
                             color=st.get("color"))
                else:
                    _set_run(pa.add_run(str(v)), cjk="微软雅黑", size=size)
        if widths:
            for r in t.rows:
                for i, w in enumerate(widths):
                    r.cells[i].width = Cm(w)
        self._tbl_borders(t, width=4)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(0)
        _set_run(sp.add_run(""), size=4)
        return t

    def page_break(self):
        self.doc.add_page_break()

    def save(self, path):
        self.doc.save(path)
